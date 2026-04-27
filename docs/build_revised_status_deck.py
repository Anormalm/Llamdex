import sys
sys.path.insert(0, '/tmp/pptx_site')
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

ROOT = Path('/home/anormalm/Llamdex')
OUT = ROOT / 'llamdex_plan1pp_status_20260427_revised.pptx'
ASSETS = ROOT / 'docs' / 'slide_assets'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

NAVY = RGBColor(10, 25, 47)
INK = RGBColor(28, 35, 49)
MUTED = RGBColor(91, 103, 120)
BG = RGBColor(247, 244, 238)
PANEL = RGBColor(255, 255, 255)
ACCENT = RGBColor(213, 106, 59)
TEAL = RGBColor(33, 120, 123)
GREEN = RGBColor(56, 143, 96)
RED = RGBColor(180, 74, 65)
GOLD = RGBColor(191, 146, 55)

FONT_HEAD = 'Aptos Display'
FONT_BODY = 'Aptos'


def add_bg(slide, section=''):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background()
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), prs.slide_height)
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    if section:
        box = slide.shapes.add_textbox(Inches(0.44), Inches(6.95), Inches(5.0), Inches(0.25))
        p = box.text_frame.paragraphs[0]
        p.text = section.upper()
        p.font.name = FONT_BODY; p.font.size = Pt(8); p.font.bold = True; p.font.color.rgb = MUTED


def add_title(slide, title, kicker=None, y=0.45):
    if kicker:
        k = slide.shapes.add_textbox(Inches(0.72), Inches(y), Inches(11.8), Inches(0.25))
        p = k.text_frame.paragraphs[0]
        p.text = kicker.upper()
        p.font.name = FONT_BODY; p.font.size = Pt(9); p.font.bold = True; p.font.color.rgb = ACCENT
        y += 0.28
    t = slide.shapes.add_textbox(Inches(0.72), Inches(y), Inches(11.8), Inches(0.70))
    p = t.text_frame.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEAD; p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = NAVY
    return y + 0.84


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_BODY; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    if align: p.alignment = align
    return box


def add_bullets(slide, bullets, x, y, w, h, size=16, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.name = FONT_BODY; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def card(slide, x, y, w, h, title=None, fill=PANEL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(226, 220, 210)
    if title:
        add_text(slide, title, x+0.22, y+0.18, w-0.44, 0.28, size=11, color=ACCENT, bold=True)
    return shape


def metric_card(slide, x, y, w, h, label, value, sub='', color=NAVY):
    card(slide, x, y, w, h)
    add_text(slide, value, x+0.18, y+0.22, w-0.36, 0.45, size=25, color=color, bold=True)
    add_text(slide, label, x+0.18, y+0.76, w-0.36, 0.35, size=12, color=INK, bold=True)
    if sub:
        add_text(slide, sub, x+0.18, y+1.12, w-0.36, 0.44, size=10, color=MUTED)


def add_table(slide, rows, x, y, w, h, font_size=10):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for c in range(len(rows[0])):
        table.columns[c].width = Inches(w / len(rows[0]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else PANEL
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_BODY; p.font.size = Pt(font_size); p.font.color.rgb = RGBColor(255,255,255) if r == 0 else INK
                p.font.bold = r == 0
    return table


def add_image_fit(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

# 1
s = prs.slides.add_slide(blank); add_bg(s, 'status')
add_text(s, 'LLAMDEX PLAN-1++', 0.78, 0.58, 5, 0.35, size=11, color=ACCENT, bold=True)
add_text(s, 'Privacy-Bounded\nMultimodal Injection', 0.78, 1.0, 7.0, 1.3, size=34, color=NAVY, bold=True)
add_text(s, 'Corrected baselines, architecture status, and the next publishable claim', 0.82, 2.55, 6.9, 0.55, size=17, color=INK)
metric_card(s, 0.82, 3.55, 2.4, 1.45, 'DINOv2 probe', '85.2%', 'corrected DTD baseline', GREEN)
metric_card(s, 3.45, 3.55, 2.4, 1.45, 'CLIP probe', '82.0%', 'corrected DTD baseline', TEAL)
metric_card(s, 6.08, 3.55, 2.4, 1.45, 'Yes/no task', '98.4%', 'router-parallel runs', NAVY)
add_text(s, 'Advisor / research group update\nApril 28, 2026', 9.25, 5.72, 3.1, 0.7, size=13, color=MUTED, align=PP_ALIGN.RIGHT)

# 2
s = prs.slides.add_slide(blank); add_bg(s, 'message'); add_title(s, 'One-slide status message')
add_bullets(s, [
    'Architecture direction is intact: private client evidence is injected into a frozen server model through small trainable adapters.',
    'Evaluation is stricter now: DTD and Oxford labels use real dataset ontologies, and CODE=label parsing is fixed.',
    'The baseline story changed: adapted representation baselines are strong and are the correct comparator.',
    'Publishable framing should be privacy-bounded adaptation and grounded behavior, not image-classification SOTA yet.'
], 0.95, 1.55, 7.0, 3.7, size=18)
card(s, 8.45, 1.55, 3.85, 3.7, 'Claim boundary')
add_bullets(s, ['Can say: auditable benchmark pipeline and working injection policies.', 'Cannot say yet: SOTA DTD classification.', 'Need next: retrain connector under corrected labels.'], 8.75, 2.10, 3.3, 2.5, size=14)

# 3
s = prs.slides.add_slide(blank); add_bg(s, 'problem'); add_title(s, 'Research question')
add_text(s, 'Can we adapt a frozen multimodal LLM without sending private raw data to the server?', 0.95, 1.35, 11.0, 0.55, size=22, color=NAVY, bold=True)
for i,(name,txt,col) in enumerate([
    ('Client', 'private image\nlocal expert\nevidence bundle', TEAL),
    ('Boundary', 'compact evidence only\nmanifested contract\nno raw image upload', ACCENT),
    ('Server', 'frozen backbone\nconnector/router\nconstrained output', NAVY),
]):
    x = 1.05 + i*4.1
    card(s, x, 2.25, 3.35, 2.35)
    add_text(s, name, x+0.22, 2.48, 2.8, 0.35, size=18, color=col, bold=True)
    add_text(s, txt, x+0.22, 3.05, 2.8, 1.05, size=16, color=INK)
add_text(s, 'Evaluation target: compare injection policies against strong local representation and API baselines.', 1.0, 5.55, 11.0, 0.45, size=17, color=MUTED)

# 4
s = prs.slides.add_slide(blank); add_bg(s, 'architecture'); add_title(s, 'System overview')
add_image_fit(s, ASSETS/'server_client_architecture.png', 1.05, 1.35, 10.9, 5.45)

# 5
s = prs.slides.add_slide(blank); add_bg(s, 'architecture'); add_title(s, 'Injection mechanisms')
add_image_fit(s, ASSETS/'injection_mechanisms.png', 0.9, 1.18, 11.6, 5.95)

# 6
s = prs.slides.add_slide(blank); add_bg(s, 'policies'); add_title(s, 'Policy variants under test')
card(s, 0.9, 1.35, 5.5, 4.85, 'Pre-FFN router')
add_bullets(s, ['Injects evidence before the feed-forward block.', 'Can improve task routing and answer selection.', 'Current runs: slightly higher fine-grained and grounded-generation accuracy.'], 1.22, 2.0, 4.8, 2.4, size=15)
card(s, 6.9, 1.35, 5.5, 4.85, 'Post-attn-layers')
add_bullets(s, ['Injects after attention across layers.', 'Better for nonempty and faithful rationales.', 'Current runs: much stronger rationale behavior.'], 7.22, 2.0, 4.8, 2.4, size=15)
add_text(s, 'Decision point: answer accuracy and rationale faithfulness are currently optimized by different injection sites.', 1.0, 6.45, 11.3, 0.4, size=16, color=ACCENT, bold=True)

# 7
s = prs.slides.add_slide(blank); add_bg(s, 'evaluation'); add_title(s, 'Evaluation protocol now has guardrails')
add_table(s, [
    ['Axis', 'Current protocol'],
    ['Fine-grained classification', 'DTD texture labels; Oxford-IIIT Pet labels'],
    ['Strict yes/no', 'Binary visual assertions with constrained decoding'],
    ['Population estimation', 'Numeric/relative estimation with MAE'],
    ['Grounded generation', 'Answer plus rationale/evidence consistency'],
], 0.9, 1.35, 7.3, 3.0, font_size=12)
card(s, 8.65, 1.35, 3.6, 3.0, 'New guardrails')
add_bullets(s, ['Real class-name ontologies.', 'Robust CODE=label parser.', 'Raw prediction audit JSONL.', 'DINOv2 and CLIP probes.'], 8.95, 1.88, 3.0, 2.0, size=13)
add_text(s, 'This is the difference between a demo benchmark and a defensible research artifact.', 1.0, 5.45, 11.0, 0.45, size=18, color=NAVY, bold=True)

# 8
s = prs.slides.add_slide(blank); add_bg(s, 'baseline'); add_title(s, 'Baseline audit: why the old result was too low')
add_bullets(s, [
    'DTD class names were placeholders, so old zero-shot prompts did not encode the real texture label space.',
    'Some CODE=label responses were parsed by the label text rather than the option code.',
    'Earlier slides mixed current rows with stale artifacts and historical API numbers.',
    'Fresh API rerun is blocked by quota, so API numbers must be labeled historical until rerun.'
], 0.95, 1.4, 7.2, 3.6, size=17)
card(s, 8.55, 1.45, 3.75, 3.25, 'Publishable correction')
add_bullets(s, ['Corrected labels.', 'Parser regression tests.', 'Auditable raw predictions.', 'Representation baselines included.'], 8.85, 2.0, 3.1, 2.0, size=14)

# 9
s = prs.slides.add_slide(blank); add_bg(s, 'baselines'); add_title(s, 'Corrected DTD baseline snapshot')
add_text(s, 'One-shot semantic-v2 protocol, 128 eval images', 0.95, 1.18, 7.5, 0.3, size=13, color=MUTED)
add_table(s, [
    ['Method', 'Acc.', 'Macro-F1', 'Reading'],
    ['DINOv2 linear probe', '85.2%', '79.2%', 'strong local baseline'],
    ['CLIP linear probe', '82.0%', '77.1%', 'strong local baseline'],
    ['Current injection checkpoint', '9.4%', 'n/a', 'stale/mismatched row'],
    ['Caption + LLM', '7.0%', '2.0%', 'weak for texture taxonomy'],
    ['LLM only', '3.1%', '0.2%', 'near random over 47 classes'],
    ['Frozen VLM prompts', '2-3%', '0-1%', 'prompting is insufficient'],
], 0.75, 1.6, 11.9, 3.7, font_size=10)
add_text(s, 'Main reading: the meaningful baseline is adapted visual representation performance, not zero-shot chat prompting.', 0.95, 5.75, 11.2, 0.45, size=17, color=ACCENT, bold=True)

# 10
s = prs.slides.add_slide(blank); add_bg(s, 'task matrix'); add_title(s, 'Task-matrix connector status')
add_table(s, [
    ['Metric', 'Pre-FFN', 'Post-attn', 'Reading'],
    ['Fine-grained acc.', '53.9%', '51.2%', 'pre-FFN slightly higher'],
    ['Strict yes/no acc.', '98.4%', '98.4%', 'both strong'],
    ['Population MAE', '0.0228', '0.0228', 'no separation'],
    ['Grounded-gen acc.', '55.5%', '53.9%', 'similar'],
    ['Rationale nonempty', '6.8%', '72.4%', 'post-attn much better'],
    ['Rationale faithful', '5.7%', '29.2%', 'post-attn much better'],
], 0.75, 1.42, 11.8, 3.65, font_size=10)
add_text(s, 'Takeaway: injection is useful across the task matrix, but the same site does not optimize classification and rationales.', 0.95, 5.62, 11.0, 0.5, size=17, color=NAVY, bold=True)

# 11
s = prs.slides.add_slide(blank); add_bg(s, 'claims'); add_title(s, 'Current claim boundary')
card(s, 0.9, 1.35, 5.55, 4.5, 'Can say now')
add_bullets(s, ['Privacy-bounded client/server evidence-injection contract is implemented.', 'Benchmark pipeline now has auditable labels, parsing, and raw predictions.', 'Pre-FFN and post-attn injection show distinct strengths across task types.'], 1.22, 1.95, 4.85, 2.7, size=15)
card(s, 6.9, 1.35, 5.55, 4.5, 'Should not say yet')
add_bullets(s, ['Not SOTA on DTD classification.', 'Not apples-to-apples against current API models until quota issue is resolved.', 'Not final connector result until retrained under corrected semantic labels.'], 7.22, 1.95, 4.85, 2.7, size=15)

# 12
s = prs.slides.add_slide(blank); add_bg(s, 'plan'); add_title(s, 'Publishable path')
add_table(s, [
    ['Step', 'Output needed', 'Why it matters'],
    ['Retrain connector', 'new injection rows', 'fair comparison to DINOv2/CLIP'],
    ['Rerun API baseline', 'current external comparator', 'avoid stale overclaiming'],
    ['Add audit appendix', 'raw predictions + confusion slices', 'defensible failure analysis'],
    ['Compare policies', 'pre-FFN vs post-attn same protocol', 'supports architecture claim'],
    ['Freeze manifest', 'configs, commits, run paths', 'reproducible report'],
], 0.75, 1.35, 11.85, 3.65, font_size=10)
add_text(s, 'Target claim: privacy-bounded connector adaptation with grounded output, not universal image-classification SOTA.', 0.95, 5.63, 11.0, 0.5, size=17, color=ACCENT, bold=True)

# 13
s = prs.slides.add_slide(blank); add_bg(s, 'discussion'); add_title(s, "Tomorrow's discussion")
add_bullets(s, [
    'Is the publishable claim better framed as privacy-bounded adaptation rather than classification SOTA?',
    'Which comparator should be mandatory: DINOv2/CLIP probes, API VLMs, or both?',
    'Should the next connector training optimize classification first, rationale faithfulness first, or a multi-objective mix?',
    'What privacy threat-model detail is needed before writing the paper section?'
], 0.95, 1.55, 8.2, 3.7, size=18)
card(s, 9.55, 1.55, 2.65, 3.7, 'Artifact anchors')
add_text(s, 'Corrected DTD baseline\nTask-matrix pre-FFN\nTask-matrix post-attn\nAPI rerun: quota-blocked', 9.85, 2.12, 2.0, 1.9, size=12, color=INK)

prs.save(OUT)
print(OUT)
